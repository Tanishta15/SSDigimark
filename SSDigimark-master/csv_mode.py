import csv
from checker import process_text
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font

INPUT_FILE = "input_answers.csv"
OUTPUT_XLSX = "output_results.xlsx"

def format_clean_feedback(result):
    highlighted_sections = []

    for entry in result["errors"]:
        highlighted = f"Incorrect: {entry['highlighted']}"
        corrected = f"Corrected: {entry['corrected']}"
        highlighted_sections.append(f"{highlighted}\n{corrected}")

    highlighted_text = "\n\n".join(highlighted_sections) if highlighted_sections else "No grammar issues detected."
    score_line = f"\n\nGrammar Score: {result['score']}/100 ({result['remark']})"

    return f"Original Answer:\n{result['original']}\n\n{highlighted_text}{score_line}".replace("\n", "\r\n")


def process_csv(input_path, output_path):
    with open(input_path, newline='', encoding='utf-8') as infile, \
         open(output_path, 'w', newline='', encoding='utf-8') as outfile, \
         open("output_results.txt", 'w', encoding='utf-8') as txtfile:

        reader = csv.DictReader(infile)
        fieldnames = ["AnswerID", "Feedback", "Corrected"]
        writer = csv.DictWriter(outfile, fieldnames=fieldnames)
        writer.writeheader()
        rows = []

        for row in reader:
            answer_id = row.get("AnswerID", "").strip()
            answer_text = row.get("AnswerText", "").strip()
            if not answer_text:
                continue

            result = process_text(answer_text)
            feedback = format_clean_feedback(result)

            output_row = {
                "AnswerID": answer_id,
                "Feedback": feedback,
                "Corrected": result["corrected"]
            }

            writer.writerow(output_row)
            rows.append(output_row)

            txtfile.write(f"AnswerID: {answer_id}\n{feedback}\n\n{'-'*80}\n")
            txtfile.write(f"\nCorrected Answer:\n{result['corrected']}\n\n\n")

    convert_csv_to_xlsx(rows, OUTPUT_XLSX)
    print(f"\nClean feedback saved to:\n - CSV: '{output_path}'\n - TXT: 'output_results.txt'\n - XLSX: '{OUTPUT_XLSX}'")


def convert_csv_to_xlsx(rows, xlsx_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Grammar Feedback"

    headers = ["AnswerID", "Feedback", "Corrected"]
    ws.append(headers)

    # Optional formatting tweaks
    bold_font = Font(bold=True)
    for col_num, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = bold_font
        ws.column_dimensions[get_column_letter(col_num)].width = 40

    for row in rows:
        ws.append([row["AnswerID"], row["Feedback"], row["Corrected"]])

    # Wrap text and align vertically
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(xlsx_path)


if __name__ == "__main__":
    process_csv(INPUT_FILE, "output_results.csv")
